# Builds templates/Plan_Proyectos_Template.xlsx, the crew lead's day plan
# that Master Schedule's "Import day plan" reads (day_plan.js). Run it again
# after the Connecteam roster changes, so the name dropdown is current:
#   python3 scripts/make_day_plan_template.py
# The row labels in column A and the "Responsabilidad | Nombre" header row
# are what day_plan.js looks for; change them there too.
import json, re, datetime, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'templates', 'Plan_Proyectos_Template.xlsx')
roster = sorted({re.sub(r'\s+', ' ', p['name']).strip() for p in json.load(open(os.path.join(ROOT, 'data', 'connecteam_roster.json')))['people']}, key=str.lower)
TODAY = datetime.date.today().isoformat()

F = 'Arial'
def font(**k): return Font(name=F, **k)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
CREAM, ORANGE, GREEN, RED, GREY, YELLOW, BLUE = 'FFF8E7', 'F4A340', 'C6EFCE', 'E06666', 'EDEDED', 'FFEB84', 'DDEBF7'
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)

NB = 10                     # project blocks per day
FIRST, LAST = 15, 64        # people rows
ROLE = lambda i: 2 + 2 * i  # B, D, F ...
NAME = lambda i: 3 + 2 * i  # C, E, G ...
LASTCOL = 1 + 2 * NB         # last visible column (U for 10 blocks)
# Hidden helpers past the blocks: last role / effective role / clean name.
HL = LASTCOL + 3; HE = HL + NB; HC = HE + NB
EFF_RANGE = f"${L(HE)}${FIRST}:${L(HE+NB-1)}${LAST}"
CLEAN_RANGE = f"${L(HC)}$6:${L(HC+NB-1)}${LAST}"
clean = lambda ref: f'TRIM(IFERROR(LEFT({ref},FIND("(",{ref})-1),{ref}))'

# Responsibility → Labor Log process, as ROLE_PROCESS in day_plan.js.
ROLES = [
  ('Cortar', 'Harvest / cut', 'harvest'),
  ('Fresh frozen', 'Fresh frozen harvest', 'fresh_frozen'),
  ('Colgar', 'Hang (wet intake)', 'intake_wet'),
  ('Pesar cajas', 'Weigh boxes (wet intake)', 'intake_wet'),
  ('Metrc', 'Metrc / tags', '(project\'s process)'),
  ('Cuartos', 'Rooms', 'other'),
  ('Descolgar', 'Take down (dry)', 'dry_check'),
  ('Destronche', 'Bucking / de-stem', 'buck'),
  ('Desoje', 'Leaf strip', 'buck'),
  ('Pesar', 'Weigh (bucking)', 'buck'),
  ('Comodín proceso', 'Floater - process', 'buck'),
  ('Mobius', 'Mobius machine', 'machine_trim'),
  ('Trim', 'Hand trim', 'hand_trim'),
  ('Smalls', 'Smalls trim', 'smalls_trim'),
  ('Prerolls', 'Pre-rolls', 'preroll_production'),
  ('Flor molida', 'Ground flower pouches', 'pouch_production'),
  ('Packing', 'Packing', '(project\'s process)'),
  ('Chofer', 'Driver', 'transport'),
  ('Regar', 'Watering', 'other'),
  ('Bomba', 'Pump', 'other'),
  ('Security', 'Security / guard', 'other'),
  ('Otro', 'Other - explain in Notas', 'other'),
]
SITES = ['Adobe', 'Airway', 'Wildcat', 'Comstock', 'Lucerne', 'Sulphur Bank', 'Highland', 'Otro']

wb = Workbook()

def plan_sheet(ws, title_note, date=None):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 24
    for i in range(NB):
        ws.column_dimensions[L(ROLE(i))].width = 16
        ws.column_dimensions[L(NAME(i))].width = 24
    for c in range(LASTCOL + 1, HC + NB):
        ws.column_dimensions[L(c)].hidden = True
    ws['A1'] = 'PLAN DEL DÍA'; ws['A1'].font = font(bold=True, size=14)
    ws['B1'] = 'Fecha / Date:'; ws['B1'].font = font(bold=True); ws['B1'].alignment = Alignment(horizontal='right')
    ws.merge_cells('C1:D1'); ws['C1'] = date; ws['C1'].number_format = 'yyyy-mm-dd'
    ws['C1'].fill = fill(CREAM); ws['C1'].font = font(bold=True, size=12); ws['C1'].border = box
    ws['F1'] = 'Total personas / people:'; ws['F1'].font = font(bold=True); ws['F1'].alignment = Alignment(horizontal='right')
    ws['G1'] = '=' + '+'.join(f'{L(ROLE(i))}13' for i in range(NB)); ws['G1'].font = font(bold=True, size=12)
    ws['I1'] = 'Nombres repetidos / Repeated:'; ws['I1'].font = font(bold=True); ws['I1'].alignment = Alignment(horizontal='right')
    ws['J1'] = f'=SUMPRODUCT(({CLEAN_RANGE}<>"")*(COUNTIF({CLEAN_RANGE},{CLEAN_RANGE})>1))'; ws['J1'].font = font(bold=True, size=12)
    ws['A2'] = title_note; ws['A2'].font = font(italic=True, color='595959')
    labels = {3: 'Proyecto / Project', 4: 'Sitio / Site', 5: 'Horario: inicio | fin', 6: 'Líder / Lead',
              7: 'Contratista 1: nombre | cuántos', 8: 'Contratista 1: inicio | fin',
              9: 'Contratista 2: nombre | cuántos', 10: 'Contratista 2: inicio | fin',
              11: 'Notas / Notes', 12: 'Nombrados (incl. líder)', 13: 'TOTAL (con contratistas)',
              14: 'PERSONAS ↓'}
    for r, t in labels.items():
        c = ws.cell(r, 1, t); c.font = font(bold=True, size=9); c.alignment = left; c.border = box; c.fill = fill(GREY)
    ws.cell(14, 1).value = 'Responsabilidad | Nombre ↓'
    for r in range(FIRST, LAST + 1):
        ws.cell(r, 1, r - FIRST + 1).font = font(size=8, color='999999')
        ws.cell(r, 1).alignment = Alignment(horizontal='right')
    ws.row_dimensions[11].height = 30
    for i in range(NB):
        a, b = ROLE(i), NAME(i)
        A, B = L(a), L(b)
        for r in (3, 4, 6, 11, 12, 13):
            ws.merge_cells(f'{A}{r}:{B}{r}')
        for r in range(3, 14):
            for c in (a, b):
                cell = ws.cell(r, c); cell.border = box; cell.font = font(); cell.alignment = center
                if r <= 11: cell.fill = fill(CREAM)
        ws[f'{A}3'].fill = fill(ORANGE); ws[f'{A}3'].font = font(bold=True, size=11)
        ws[f'{A}6'].fill = fill(GREEN); ws[f'{A}6'].font = font(bold=True)
        ws[f'{A}12'] = f'=({A}6<>"")+COUNTA({B}{FIRST}:{B}{LAST})'
        ws[f'{A}13'] = f'={A}12+N({B}7)+N({B}9)'; ws[f'{A}13'].font = font(bold=True)
        for c, t in ((a, 'Responsabilidad'), (b, 'Nombre')):
            h = ws.cell(14, c, t); h.font = font(bold=True, size=9, color='FFFFFF'); h.fill = fill('595959'); h.alignment = center; h.border = box
        for r in range(FIRST, LAST + 1):
            for c in (a, b):
                cell = ws.cell(r, c); cell.fill = fill(CREAM); cell.border = box; cell.font = font(size=10)
            # helpers: carried-down responsibility, effective responsibility, name without (notes)
            ws.cell(r, HL + i, f'=IF({A}{r}<>"",{clean(f"{A}{r}")},"")' if r == FIRST else f'=IF({A}{r}<>"",{clean(f"{A}{r}")},{L(HL+i)}{r-1})')
            ws.cell(r, HE + i, f'=IF({B}{r}="","",IF({L(HL+i)}{r}="","(sin responsabilidad)",{L(HL+i)}{r}))')
            ws.cell(r, HC + i, f'=IF({B}{r}="","",{clean(f"{B}{r}")})')
        ws.cell(6, HC + i, f'=IF({A}6="","",{clean(f"{A}6")})')
        # responsibility cells turn red once filled
        ws.conditional_formatting.add(f'{A}{FIRST}:{A}{LAST}', FormulaRule(formula=[f'{A}{FIRST}<>""'], fill=fill(RED), font=Font(name=F, bold=True, color='FFFFFF')))
        # repeated names (lead included; text in parentheses ignored)
        for rng, top in ((f'{B}{FIRST}:{B}{LAST}', f'{B}{FIRST}'), (f'{A}6', f'{A}6')):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({top}<>"",COUNTIF({CLEAN_RANGE},{clean(top)})>1)'], fill=fill(YELLOW)))
    # dropdowns (warning only: anything can still be typed)
    def dv(src, ranges, msg):
        d = DataValidation(type='list', formula1=src, allow_blank=True, showErrorMessage=True, errorStyle='warning',
                           errorTitle='No está en la lista', error=msg)
        ws.add_data_validation(d)
        for rg in ranges: d.add(rg)
    dv(f'Listas!$A$2:$A$60', [f'{L(ROLE(i))}{FIRST}:{L(ROLE(i))}{LAST}' for i in range(NB)],
       'Esta responsabilidad no está en Listas. ¿Agregarla? / Not on the Listas tab.')
    dv(f'Personal!$A$2:$A$400', [f'{L(NAME(i))}{FIRST}:{L(NAME(i))}{LAST}' for i in range(NB)] + [f'{L(ROLE(i))}6' for i in range(NB)],
       'Este nombre no está en Personal (Connecteam). Revisa cómo se escribe. / Name not on the roster.')
    dv(f'Listas!$E$2:$E$12', [f'{L(ROLE(i))}4' for i in range(NB)], 'Sitio nuevo: escríbelo en Notas también.')
    ws.freeze_panes = 'B4'
    ws.sheet_properties.tabColor = 'F4A340'

plan = wb.active; plan.title = 'Plan del Día'
plan_sheet(plan, 'Escribe solo en las celdas crema. Un bloque (2 columnas) por proyecto. Cada persona en su fila con su responsabilidad al lado; responsabilidad en blanco = la misma de arriba.')

# ── Ejemplo ─────────────────────────────────────────────────────────────────
ex = wb.create_sheet('Ejemplo')
plan_sheet(ex, 'EJEMPLO (de la foto del 19/10, con nombres como están en Connecteam). No se conecta con Resumen. / Sample only.', datetime.date(2026, 10, 19))
EXAMPLE = [
  dict(p='Bucking', s='Adobe', t=('7am', '3:30pm'), lead='', people=[
    ('Destronche', ['Isabel Hernandez', 'Araceli Martinez', 'Gandy Cardenas', 'Ramona Ortiz', 'Dalett Castro', 'Maria C Pille', 'Elena Garnica']),
    ('Desoje', ['Rosa Baeza', 'Maria Herlinda Zavala', 'Raquel Garcia', 'Alma Garcia', 'Yuliana Rivas Garcia', 'Sr. Juan Ninche', 'Flor Sierra', 'Maria Elena Mendoza', 'Rosa Navarro', 'Brenda Olivas', 'Elisa Cervantes']),
    ('Pesar', ['Ernestina Garcia']), ('Comodín proceso', ['David Cruz']), ('Mobius', ['Alicia Morales'])]),
  dict(p='Prerolls', s='Adobe', t=('7am', '3:30pm'), lead='Jesus Ayerin Sicairos', people=[
    ('Prerolls', ['Vianey Zazueta', 'Araceli Muniz', 'Sirenia Zurita', 'Ana Reyes']), ('Packing', ['Maria Toledo'])]),
  dict(p='Trimming', s='Adobe', t=('7am', '3:30pm'), lead='', c1=('Juan', 15), people=[]),
  dict(p='Harvest Intake', s='Adobe', t=('5am', '2pm'), lead='', c1=('Juan', 8), c2=('Juan', 15, '5:30pm', '2am'),
       notes='Colgar a las 5:30pm: contratista 2', people=[
    ('Colgar', ['Jorge Santiago Calvo', 'Jesus Medina', 'Jenri Yobani Lopez Lopez', 'Carlos Tentle', 'Roberto Castro', 'Juan Carlos Lopez', 'David Martinez Juarez', 'Alex Palacios', 'Eduardo Leon']),
    ('Cuartos', ['Kevin Sicairos'])]),
  dict(p='Harvest', s='Comstock', t=('4am', '12:30pm'), lead='Yobani Vega', c1=('Juan', 10), people=[
    ('Cortar', ['Jose Guadalupe Lopez', 'Diego Lopez', 'Jose Fredi Aguilar L.', 'Ismael Cruz', 'Lisandro Molina', 'Jonathan Garcia Perez', 'Alvaro Cruz Lopez', 'Edi Yunue Vega Bustamante', 'Gilberto Diaz']),
    ('Chofer', ['Jorge Rojas (sale 5am)', 'Ernesto Daniel Sicairos (sale 4am)', 'Antonio Rivas (sale 3am)', 'Antonio Sonato (sale 6am)'])]),
  dict(p='Watering', s='Wildcat', t=('7am', '3:30pm'), lead='', people=[('Regar', ['Isai Cruz Lopez'])]),
  dict(p='Fresh Frozen', s='Sulphur Bank', t=('7am', '3:30pm'), lead='Juan J. Parra', c1=('Juan', 20), people=[
    ('Fresh frozen', ['Anastasia Santiago', 'Carla Gomez Perez', 'Aurelia Rosales'])]),
]
for i, e in enumerate(EXAMPLE):
    a, b = ROLE(i), NAME(i)
    ex.cell(3, a, e['p']); ex.cell(4, a, e['s']); ex.cell(5, a, e['t'][0]); ex.cell(5, b, e['t'][1])
    if e['lead']: ex.cell(6, a, e['lead'])
    for key, row in (('c1', 7), ('c2', 9)):
        if key in e:
            ex.cell(row, a, e[key][0]); ex.cell(row, b, e[key][1])
            if len(e[key]) > 2: ex.cell(row + 1, a, e[key][2]); ex.cell(row + 1, b, e[key][3])
    if e.get('notes'): ex.cell(11, a, e['notes'])
    r = FIRST
    for role, names in e['people']:
        for k, n in enumerate(names):
            if k == 0: ex.cell(r, a, role)
            ex.cell(r, b, n); r += 1
ex.sheet_properties.tabColor = 'BFBFBF'

# ── Resumen ─────────────────────────────────────────────────────────────────
rs = wb.create_sheet('Resumen')
rs.sheet_view.showGridLines = False
P = "'Plan del Día'!"
rs['A1'] = 'RESUMEN / SUMMARY'; rs['A1'].font = font(bold=True, size=14)
rs['A2'] = "Se llena solo desde 'Plan del Día'. No escribir aquí. / Auto-filled."; rs['A2'].font = font(italic=True, color='595959')
rs['H1'] = 'Fecha:'; rs['H1'].font = font(bold=True); rs['I1'] = f'=IF({P}C1="","",{P}C1)'; rs['I1'].number_format = 'yyyy-mm-dd'
hdr = ['Proyecto', 'Sitio', 'Inicio', 'Fin', 'Líder', 'Nombrados', 'Contratistas', 'Contratistas #', 'TOTAL', 'Notas']
widths = [20, 14, 9, 9, 24, 11, 24, 13, 9, 40]
for j, (h, w) in enumerate(zip(hdr, widths), 1):
    c = rs.cell(4, j, h); c.font = font(bold=True, color='FFFFFF'); c.fill = fill('595959'); c.alignment = center; c.border = box
    rs.column_dimensions[L(j)].width = w
for i in range(NB):
    r = 5 + i; A, B = L(ROLE(i)), L(NAME(i))
    g = lambda ref: f'=IF({P}{ref}="","",{P}{ref})'
    vals = [g(f'{A}3'), g(f'{A}4'), g(f'{A}5'), g(f'{B}5'), g(f'{A}6'),
            f'=IF({P}{A}13=0,"",{P}{A}12)',
            f'=TRIM(IF({P}{A}7="","",{P}{A}7&" ×"&{P}{B}7)&IF({P}{A}9="",""," · "&{P}{A}9&" ×"&{P}{B}9&" "&{P}{A}10))',
            f'=IF(N({P}{B}7)+N({P}{B}9)=0,"",N({P}{B}7)+N({P}{B}9))',
            f'=IF({P}{A}13=0,"",{P}{A}13)', g(f'{A}11')]
    for j, v in enumerate(vals, 1):
        c = rs.cell(r, j, v); c.font = font(); c.border = box; c.alignment = left if j in (1, 5, 7, 10) else center
tr = 5 + NB
rs.cell(tr, 1, 'TOTAL').font = font(bold=True)
for j, col in ((6, 'F'), (8, 'H'), (9, 'I')):
    c = rs.cell(tr, j, f'=SUM({col}5:{col}{tr-1})'); c.font = font(bold=True); c.alignment = center; c.border = box
br = tr + 3
rs.cell(br, 1, 'PERSONAS POR RESPONSABILIDAD / PEOPLE BY RESPONSIBILITY').font = font(bold=True, size=12)
rs.cell(br + 1, 1, 'Responsabilidad'); rs.cell(br + 1, 2, 'Personas')
for j in (1, 2):
    c = rs.cell(br + 1, j); c.font = font(bold=True, color='FFFFFF'); c.fill = fill('595959'); c.border = box
first = br + 2
NL = 40
for k in range(NL):
    r = first + k
    rs.cell(r, 1, f'=IF(Listas!A{k+2}="","",Listas!A{k+2})')
    rs.cell(r, 2, f'=IF(A{r}="","",COUNTIF({P}{EFF_RANGE},A{r}))')
last = first + NL - 1
r = last + 1
rs.cell(r, 1, '(sin responsabilidad)'); rs.cell(r, 2, f'=COUNTIF({P}{EFF_RANGE},"(sin responsabilidad)")')
rs.cell(r + 1, 1, 'Otra (no está en Listas)')
rs.cell(r + 1, 2, f'=COUNTIF({P}{EFF_RANGE},"?*")-SUM(B{first}:B{r})')
for rr in range(first, r + 2):
    for j in (1, 2):
        c = rs.cell(rr, j); c.font = font(); c.border = box
rs.cell(r + 2, 1, 'Sin contar líderes ni contratistas. "(sin responsabilidad)" = nombre sin responsabilidad arriba: corregir.').font = font(italic=True, size=9, color='595959')
rs.conditional_formatting.add(f'A{r}:B{r+1}', FormulaRule(formula=[f'$B{r}>0'], fill=fill(YELLOW)))
rs.sheet_properties.tabColor = '4472C4'

# ── Cómo usar ───────────────────────────────────────────────────────────────
hw = wb.create_sheet('Cómo usar')
hw.sheet_view.showGridLines = False
hw.column_dimensions['A'].width = 26; hw.column_dimensions['B'].width = 110
hw['A1'] = 'CÓMO USAR / HOW TO USE'; hw['A1'].font = font(bold=True, size=14)
rows = [
  ('Idea', 'Cada bloque de 2 columnas es un PROYECTO del día (como en la Master Schedule). Abajo, cada persona en su fila con su RESPONSABILIDAD al lado. / Each 2-column block is one project; each person on a row with their responsibility next to them.'),
  ('Cada día / Each day', "1) Duplica 'Plan del Día' para guardarlo (clic derecho > Duplicar) y ponle la fecha. 2) Borra las celdas crema y llena el nuevo día. 3) En la app: Plan Maestro → Proyectos → ⇪ Importar plan del día, y sube el archivo (o pega las celdas). / Duplicate to archive, clear, fill, then Master Schedule → Projects → Import day plan."),
  ('Proyecto (naranja)', "Nombre corto del trabajo: Bucking, Prerolls, Trimming, Harvest Intake, Harvest, Watering, Fresh Frozen, Dry Take Down… La fecha se agrega sola. Mismo sitio con dos trabajos distintos = dos bloques. / Short job name; date added automatically."),
  ('Sitio', 'Menú: Adobe, Airway, Wildcat, Comstock, Lucerne, Sulphur Bank, Highland, Otro (si es Otro, escribe el lugar en Notas).'),
  ('Horario', 'Inicio a la izquierda, fin a la derecha. Texto libre: 7am, 3:30pm, 16:00. Si no sabes el fin, déjalo en blanco (se usan 8 horas).'),
  ('Líder (verde)', 'Una persona. Cuenta en el total: no la repitas abajo.'),
  ('Contratistas', 'Nombre del contratista (Juan) y CUÁNTOS (solo número). La fila de abajo es su horario solo si es distinto al del proyecto. Contratista 2 = otro turno (ej. colgar a las 5:30pm). No los pongas en la lista de personas.'),
  ('PERSONAS', 'Una persona por fila (hasta 50). Columna izquierda = RESPONSABILIDAD (se pinta roja), derecha = NOMBRE. Escribe la responsabilidad solo en la primera persona del grupo: en blanco = la misma de arriba. Usa el menú para que el nombre quede igual que en Connecteam.'),
  ('Responsabilidad primero', 'IMPORTANTE: la responsabilidad va en la MISMA fila que la primera persona del grupo, nunca en una fila sola encima o debajo. / The responsibility goes on the same row as the first person of its group.'),
  ('Notas entre paréntesis', "Lo que va entre ( ) es una nota: 'Jorge Rojas (sale 5am)', 'Colgar (5:30pm-2am)'. Una hora en la responsabilidad = esa responsabilidad tiene otro horario. No cuenta al comparar nombres."),
  ('Nombres repetidos (amarillo)', 'Si alguien aparece dos veces en el día (líder incluido) se pinta amarillo. Si de verdad trabaja en dos proyectos, pon la hora en paréntesis en los dos.'),
  ('Menús', "Si escribes algo que no está en el menú sale un aviso; puedes seguir. Nombres nuevos: agrégalos en 'Personal'. Responsabilidades nuevas: agrégalas en 'Listas'."),
  ('Resumen', 'Una fila por proyecto y personas por responsabilidad. Si sale "(sin responsabilidad)" hay un nombre sin responsabilidad arriba.'),
  ('Personal', f'Lista de Connecteam al {TODAY} (para el menú de nombres). Si un nombre está dos veces (dos "Maria Mendoza"), al importar se elige cuál.'),
  ('Columnas ocultas', f'Las columnas después de {L(LASTCOL)} son cálculos. No borrar.'),
]
for k, (a, b) in enumerate(rows, 3):
    hw.cell(k, 1, a).font = font(bold=True); hw.cell(k, 2, b).font = font()
    hw.cell(k, 1).alignment = Alignment(vertical='top', wrap_text=True); hw.cell(k, 2).alignment = Alignment(vertical='top', wrap_text=True)
k = len(rows) + 4
hw.cell(k, 1, 'Colores / Colors').font = font(bold=True)
for j, (t, c) in enumerate([('Escribir aquí / Type here', CREAM), ('Proyecto / Project', ORANGE), ('Líder / Lead', GREEN), ('Responsabilidad', RED), ('Nombre repetido / Repeated', YELLOW)]):
    cell = hw.cell(k + 1 + j, 1, t); cell.fill = fill(c); cell.font = font(); cell.border = box

# ── Listas ──────────────────────────────────────────────────────────────────
ls = wb.create_sheet('Listas')
for j, (h, w) in enumerate([('Responsabilidad', 20), ('English', 26), ('Proceso (app)', 22), ('', 3), ('Sitio / Site', 16)], 1):
    c = ls.cell(1, j, h or None); ls.column_dimensions[L(j)].width = w
    if h: c.font = font(bold=True, color='FFFFFF'); c.fill = fill('595959')
for k, (a, b, c) in enumerate(ROLES, 2):
    ls.cell(k, 1, a).fill = fill(CREAM)
    for j, v in enumerate((a, b, c), 1): ls.cell(k, j).value = v; ls.cell(k, j).font = font()
for k in range(len(ROLES) + 2, 61): ls.cell(k, 1).fill = fill(CREAM)
for k, s in enumerate(SITES, 2): ls.cell(k, 5, s).font = font()
ls['G1'] = 'Agrega responsabilidades nuevas abajo en la columna A (crema). Proceso (app) = el proceso de la Labor Log; Claude lo usa al crear los proyectos.'
ls['G1'].font = font(italic=True, color='595959')

# ── Personal ────────────────────────────────────────────────────────────────
pe = wb.create_sheet('Personal')
pe['A1'] = 'Nombre (Connecteam)'; pe['A1'].font = font(bold=True, color='FFFFFF'); pe['A1'].fill = fill('595959')
pe.column_dimensions['A'].width = 32
for k, n in enumerate(roster, 2): pe.cell(k, 1, n).font = font()
pe['C1'] = f'Lista de Connecteam al {TODAY} ({len(roster)} nombres). Agrega nombres nuevos al final.'; pe['C1'].font = font(italic=True, color='595959')

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and (c.font is None or c.font.name != F):
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, size=c.font.size, color=c.font.color)
wb.save(OUT)
print(OUT)
